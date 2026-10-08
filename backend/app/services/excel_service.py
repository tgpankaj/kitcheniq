"""
KitchenIQ — Excel Export Service
Multi-sheet workbook with KPIs, orders, dishes, customers.
"""
from io import BytesIO
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from db.connection import execute_query, execute_one


BRAND_FILL = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(bold=True, size=14, color="0F172A")
BOLD = Font(bold=True)


def _f(v):
    try:
        return float(v or 0)
    except (ValueError, TypeError):
        return 0.0


def _i(v):
    try:
        return int(v or 0)
    except (ValueError, TypeError):
        return 0


def _style_header(ws, row_num, num_cols):
    """Style a header row."""
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=row_num, column=col)
        cell.fill = BRAND_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _auto_width(ws, min_width=12, max_width=40):
    """Auto-fit column widths."""
    for col_cells in ws.columns:
        letter = get_column_letter(col_cells[0].column)
        max_len = 0
        for cell in col_cells:
            try:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            except:
                pass
        ws.column_dimensions[letter].width = min(max(min_width, max_len + 2), max_width)


def build_workbook(kitchen_id, days=30):
    """Build a multi-sheet Excel workbook. Returns bytes."""
    wb = Workbook()

    # ============================================================
    # SHEET 1: Summary
    # ============================================================
    ws = wb.active
    ws.title = "Summary"

    ws['A1'] = "KitchenIQ Business Report"
    ws['A1'].font = Font(bold=True, size=18, color="2563EB")
    ws['A2'] = f"Period: Last {days} days"
    ws['A2'].font = Font(italic=True, size=10, color="64748B")
    ws['A3'] = f"Generated: {datetime.now().strftime('%d %b %Y, %H:%M')}"
    ws['A3'].font = Font(italic=True, size=10, color="64748B")

    row = 5
    ws.cell(row=row, column=1, value="Metric").font = HEADER_FONT
    ws.cell(row=row, column=1).fill = BRAND_FILL
    ws.cell(row=row, column=2, value="Value").font = HEADER_FONT
    ws.cell(row=row, column=2).fill = BRAND_FILL
    row += 1

    summary = execute_one(
        """
        SELECT
            COUNT(*) FILTER (WHERE status != 'cancelled') AS orders,
            COUNT(*) FILTER (WHERE status = 'cancelled') AS cancelled,
            COALESCE(SUM(total_amount) FILTER (WHERE status != 'cancelled'), 0) AS revenue,
            COALESCE(SUM(discount) FILTER (WHERE status != 'cancelled'), 0) AS discounts,
            COALESCE(SUM(platform_fee) FILTER (WHERE status != 'cancelled'), 0) AS fees,
            COALESCE(SUM(food_cost) FILTER (WHERE status != 'cancelled'), 0) AS food_cost,
            COALESCE(SUM(profit) FILTER (WHERE status != 'cancelled'), 0) AS profit,
            COALESCE(AVG(total_amount) FILTER (WHERE status != 'cancelled'), 0) AS aov
        FROM orders
        WHERE kitchen_id = %s
          AND order_date >= CURRENT_DATE - INTERVAL '%s days'
        """,
        (kitchen_id, days),
    )

    metrics = [
        ("Total Orders", _i(summary['orders'])),
        ("Cancelled Orders", _i(summary['cancelled'])),
        ("Gross Revenue", f"₹{_f(summary['revenue']):,.2f}"),
        ("Discounts", f"₹{_f(summary['discounts']):,.2f}"),
        ("Platform Fees", f"₹{_f(summary['fees']):,.2f}"),
        ("Food Cost", f"₹{_f(summary['food_cost']):,.2f}"),
        ("Net Profit", f"₹{_f(summary['profit']):,.2f}"),
        ("Avg Order Value", f"₹{_f(summary['aov']):,.2f}"),
    ]

    for label, value in metrics:
        ws.cell(row=row, column=1, value=label).font = BOLD
        ws.cell(row=row, column=2, value=value)
        row += 1

    _auto_width(ws)

    # ============================================================
    # SHEET 2: Orders
    # ============================================================
    ws = wb.create_sheet("Orders")

    headers = ['Order ID', 'External ID', 'Platform', 'Date', 'Time',
               'Customer', 'Phone', 'Amount', 'Discount', 'Fee',
               'Food Cost', 'Net Revenue', 'Profit', 'Status']
    ws.append(headers)
    _style_header(ws, 1, len(headers))

    orders = execute_query(
        """
        SELECT id, external_order_id, platform, order_date::text AS date,
               order_time::text AS time, customer_name, customer_phone,
               total_amount, discount, platform_fee, food_cost,
               net_revenue, profit, status
        FROM orders
        WHERE kitchen_id = %s
          AND order_date >= CURRENT_DATE - INTERVAL '%s days'
        ORDER BY order_date DESC, id DESC
        """,
        (kitchen_id, days),
    )

    for o in orders:
        ws.append([
            f"#ORD{o['id']:04d}",
            o['external_order_id'] or '',
            (o['platform'] or '').capitalize(),
            o['date'],
            (o['time'] or '')[:5],
            o['customer_name'] or 'Guest',
            o['customer_phone'] or '',
            _f(o['total_amount']),
            _f(o['discount']),
            _f(o['platform_fee']),
            _f(o['food_cost']),
            _f(o['net_revenue']),
            _f(o['profit']),
            (o['status'] or '').capitalize(),
        ])

    _auto_width(ws)

    # ============================================================
    # SHEET 3: Dish Performance
    # ============================================================
    ws = wb.create_sheet("Dishes")

    headers = ['Dish', 'Orders', 'Quantity Sold', 'Revenue', 'Avg Price']
    ws.append(headers)
    _style_header(ws, 1, len(headers))

    dishes = execute_query(
        """
        SELECT oi.item_name AS name,
               COUNT(DISTINCT oi.order_id) AS orders,
               COALESCE(SUM(oi.quantity), 0) AS qty,
               COALESCE(SUM(oi.line_total), 0) AS revenue,
               COALESCE(AVG(oi.unit_price), 0) AS avg_price
        FROM order_items oi
        JOIN orders o ON o.id = oi.order_id
        WHERE o.kitchen_id = %s
          AND o.status != 'cancelled'
          AND o.order_date >= CURRENT_DATE - INTERVAL '%s days'
        GROUP BY oi.item_name
        ORDER BY revenue DESC
        """,
        (kitchen_id, days),
    )

    for d in dishes:
        ws.append([
            d['name'],
            _i(d['orders']),
            _i(d['qty']),
            _f(d['revenue']),
            _f(d['avg_price']),
        ])

    _auto_width(ws)

    # ============================================================
    # SHEET 4: Customers
    # ============================================================
    ws = wb.create_sheet("Customers")

    headers = ['Phone', 'Name', 'Orders', 'Total Spent', 'Avg Order']
    ws.append(headers)
    _style_header(ws, 1, len(headers))

    customers = execute_query(
        """
        SELECT customer_phone AS phone,
               COALESCE(MAX(customer_name), '') AS name,
               COUNT(*) AS freq,
               COALESCE(SUM(total_amount), 0) AS total
        FROM orders
        WHERE kitchen_id = %s
          AND customer_phone IS NOT NULL
          AND status != 'cancelled'
          AND order_date >= CURRENT_DATE - INTERVAL '%s days'
        GROUP BY customer_phone
        ORDER BY total DESC
        LIMIT 200
        """,
        (kitchen_id, days),
    )

    for c in customers:
        freq = _i(c['freq'])
        ws.append([
            c['phone'],
            c['name'] or '',
            freq,
            _f(c['total']),
            round(_f(c['total']) / max(freq, 1), 2),
        ])

    _auto_width(ws)

    # Save
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()